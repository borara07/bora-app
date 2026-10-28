# -*- coding: utf-8 -*-
"""
구글 시트 '출결관리 수정본' 의 '재원생' 탭을 읽어
수파베이스 sync_roster_from_sheet 함수에 넘길 명단(JSON)을 만듭니다.

쓰는 법
    python3 시트-명단-읽기.py  받은파일.xlsx  명단.json

시트 모양 (2026년 10월 기준)
    1줄          머리줄 (이름 · 학부모 연락처 · 학생 연락처 · 학교 · 학년 · …)
    반 이름 줄   이름 칸에만 '고1 토' 처럼 반 이름이 있고 나머지는 빈 줄
    학생 줄      그 아래로 학생이 이어집니다. 다음 반 이름 줄이 나올 때까지 같은 반

읽는 규칙
    - '학년' 칸에는 기수(24 · 25 · 26 …)가 적혀 있습니다. 앱 메모는 '고1 토 · 26기' 처럼 만듭니다.
      기수를 시트에서 읽으므로 해가 바뀌어 학년이 올라가도 이 파일은 고칠 것이 없습니다.
      한 반 안에서 학년 칸이 빈 학생이 있으면 그 반에서 가장 많이 쓴 기수를 씁니다.
    - 반 이름이 '중' 으로 시작하면 중등부, 나머지는 고등부입니다.
    - 동명이인을 가르려고 붙여 둔 꼬리표(김도윤A · 김도윤B)는 뗍니다.
      학생은 로그인할 때 본명을 치고, 둘은 학부모 연락처 뒷 4자리로 갈립니다.
    - 학부모 연락처가 빈 학생은 넣지 않고 화면에 알려 줍니다 (뒷 4자리가 학생을 가리는 기준입니다).

이 파일은 명단을 만들기만 합니다. 앱 명단을 고치는 것은 수파베이스 함수가 합니다.
"""
import json
import re
import sys
from collections import Counter

import openpyxl

SHEET = '재원생'


def text(cell):
    return '' if cell is None else str(cell).strip()


def digits(s):
    return re.sub(r'\D', '', s or '')


def cohort(cell):
    """학년 칸의 '26' · '26.0' 을 26 으로 읽습니다. 못 읽으면 None"""
    m = re.match(r'^\s*(\d{2})(?:\.0+)?\s*$', text(cell))
    return int(m.group(1)) if m else None


def read(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit("'%s' 탭이 없습니다. 시트 이름: %s" % (SHEET, wb.sheetnames))
    ws = wb[SHEET]

    rows, klass, skipped, untagged = [], '', [], []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        row = list(row) + [None] * 6
        name, parent, student, school = (text(row[0]), text(row[1]),
                                         text(row[2]), text(row[3]))
        if i == 1 or not name or name == '이름':
            continue
        if not parent and not student and not school:
            klass = re.sub(r'\s+', ' ', name)        # 반 이름 줄
            continue
        if len(digits(parent)) < 4:
            skipped.append('%d줄 %s (%s) — 학부모 연락처 없음' % (i, name, klass))
            continue
        m = re.match(r'^(.+?)\s*([A-Za-z])$', name)  # 김도윤A → 김도윤
        if m:
            untagged.append(name)
            name = m.group(1).strip()
        rows.append({'name': name, 'school': school, 'klass': klass,
                     'cohort': cohort(row[4]),
                     'parent_phone': parent, 'student_phone': student})

    # 반마다 가장 많이 쓴 기수 (학년 칸이 빈 학생에게 씁니다)
    usual = {}
    for k in set(r['klass'] for r in rows):
        found = Counter(r['cohort'] for r in rows if r['klass'] == k and r['cohort'])
        usual[k] = found.most_common(1)[0][0] if found else None

    out = []
    for r in rows:
        g = r['cohort'] or usual.get(r['klass'])
        memo = ('%s · %d기' % (r['klass'], g)) if (r['klass'] and g) else r['klass']
        out.append({
            'name': r['name'],
            'school': r['school'],
            'memo': memo,
            'group': '중등부' if r['klass'].startswith('중') else '고등부',
            'parent_phone': r['parent_phone'],
            'student_phone': r['student_phone'],
        })
    return out, skipped, untagged


def main():
    if len(sys.argv) != 3:
        raise SystemExit('쓰는 법: python3 시트-명단-읽기.py 받은파일.xlsx 명단.json')
    out, skipped, untagged = read(sys.argv[1])

    # 시트 안에서 같은 사람(이름 + 뒷 4자리)이 두 번 나오면 멈춥니다
    keys = Counter((re.sub(r'\s', '', o['name']).lower(), digits(o['parent_phone'])[-4:]) for o in out)
    twice = ['%s(%s)' % k for k, n in keys.items() if n > 1]
    if twice:
        raise SystemExit('시트에 같은 학생이 두 번 있습니다: ' + ', '.join(twice))

    with open(sys.argv[2], 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, separators=(',', ':'))

    print('학생 %d명' % len(out))
    for memo, n in Counter(o['memo'] for o in out).items():
        print('  %-18s %d' % (memo, n))
    if untagged:
        print('꼬리표를 뗀 이름:', ', '.join(untagged))
    if skipped:
        print('넣지 않은 줄:')
        for s in skipped:
            print('  ' + s)


if __name__ == '__main__':
    main()
