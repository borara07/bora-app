# 모의고사·진단평가 성적표 엑셀을 만드는 스크립트입니다.
# (선생님은 이 파일을 쓰지 않습니다. 만들어진 엑셀·PDF만 쓰시면 됩니다)
# -*- coding: utf-8 -*-
"""보라! 국어전문학원 성적표 만들기

   회차 내용(정답·배점·영역·등급컷)은 앱의 exams.js 에서 그대로 읽어 옵니다.
   선택과목이 있는 회차(고3 모의고사)와 없는 회차(중등 진단평가)를 모두 만듭니다.

   쓰는 법:
       python3 성적표-만들기.py "34차 심화모의고사" 학생답안.json 성적표.xlsx

   학생답안.json 은 이런 모양입니다.
       [ {"name":"정예원", "school":"배화여고", "choice":"화법과 작문",
          "ans":[4,4,5, ... 45개]}, ... ]
   선택과목이 없는 회차는 choice 를 적지 않아도 됩니다."""
import datetime, os, sys, json, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.pagebreak import Break
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (ParagraphProperties, CharacterProperties,
                                   RichTextProperties, Paragraph)
from openpyxl.worksheet.properties import PageSetupProperties

# ---------- 회차 내용 (exams.js 에서 읽어 옵니다) ----------
AREA_ROWS = 11          # 영역 표에 들어가는 줄 수 (남으면 비워 둡니다)

def read_exams(path):
    """exams.js 를 그대로 읽어 회차 목록을 가져옵니다."""
    js = ("const fs=require('fs'),vm=require('vm');const c={};vm.createContext(c);"
          "vm.runInContext(fs.readFileSync(%r,'utf8'),c);"
          "console.log(JSON.stringify({exams:c.EXAMS,count:c.EXAM_QUESTION_COUNT,"
          "common:c.EXAM_COMMON_COUNT,choices:c.EXAM_CHOICES}));" % path)
    return json.loads(subprocess.check_output(['node', '-e', js]).decode('utf-8'))

def digits(t): return [int(c) for c in str(t or '') if c.isdigit()]

class Round(object):
    """한 회차의 정답·배점·영역·등급컷을 선택과목별로 꺼내 줍니다."""
    def __init__(self, exam, book):
        self.title = exam['title']
        self.who   = exam.get('grade') or ''
        self.date  = exam.get('date') or ''
        self.count = exam.get('count') or book['count'] or 45
        self.all_choices = book['choices'] or []
        self.choices = [c for c in self.all_choices if (exam.get('answers') or {}).get(c)]
        self.exam = exam
    def key(self, choice):
        a = self.exam['answers']
        return digits(a.get('공통')) + (digits(a.get(choice)) if choice else [])
    def point(self, choice):
        p = self.exam['points']
        return digits(p.get('공통')) + (digits(p.get(choice)) if choice else [])
    def cuts(self, choice):
        c = self.exam['cuts']
        return [int(x) for x in str(c.get(choice) or c.get('공통') or '').split('-') if x.strip()]
    def areas(self, choice):
        out = []
        for a in self.exam['areas']:
            if a[4] == '공통' or a[4] == choice:
                out.append((a[0], a[1], int(a[2]), int(a[3])))
        return out

# ---------- 색·글꼴 (30차 성적표와 같게) ----------
# 색은 선생님이 쓰시던 성적표에서 그대로 뽑아 왔습니다.
# 형광빛 보라를 빼고 차분한 청보라 한 가지로 맞췄습니다.
BAND  = "D2D2EA"   # 제목 띠 (차분한 청보라)
TAB   = "E6E0ED"   # 이름표 (1. 학생정보 …)
TH    = "E6E0ED"   # 표 머리 바탕
SUM_  = "E6E0ED"   # 합계 줄
TD    = "DBEEF3"   # 값 줄 바탕 (옅은 하늘빛)
EDGE  = "D9D9DF"   # 표 테두리
TITLE_INK = "3F3F46"  # 제목·이름표 글자 (진회색)
THINK = "1F2937"   # 표 머리 글자 (검정)
THON  = "C00000"   # '득점' 머리글만 붉은 글씨
PURPLE= "7F62A2"   # 점수·막대그래프 (차분한 보라)
GRAY  = "6B7280"
MINT, MINT_INK = "B8E6E0", "0F4F49"
VAL = "1F2937"        # 표 안의 값은 검은색
FAM = "나눔스퀘어라운드"

def F(sz=11, b=False, c="1F2937"): return Font(name=FAM, size=sz, bold=b, color=c)
def FILL(c): return PatternFill("solid", fgColor=c)
_edge = Side(style="thin", color=EDGE)
SOFT  = Border(left=_edge, right=_edge, top=_edge, bottom=_edge)
C = Alignment(horizontal="center", vertical="center")
L = Alignment(horizontal="left",   vertical="center")
LW= Alignment(horizontal="left",   vertical="center", wrap_text=True)

def put(ws, r, c, v, font=None, fill=None, align=None, border=None, fmt=None):
    cell = ws.cell(r, c, v)
    if font: cell.font = font
    if fill: cell.fill = FILL(fill)
    if align: cell.alignment = align
    if border: cell.border = border
    if fmt: cell.number_format = fmt
    return cell

def merge(ws, r1, c1, r2, c2): ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

def tab(ws, r, c1, c2, text):
    put(ws, r, c1, text, F(12, True, TITLE_INK), fill=TAB, align=L)
    merge(ws, r, c1, r, c2); ws.row_dimensions[r].height = 21

def head(ws, r, groups, redcol=None):
    for t, c1, c2 in groups:
        cell = put(ws, r, c1, t, F(11, True, THINK), fill=TH, align=C, border=SOFT)
        if redcol is not None and c1 == redcol: cell.font = F(11, True, THON)
        if c2 > c1: merge(ws, r, c1, r, c2)
        for c in range(c1, c2 + 1): ws.cell(r, c).border = SOFT
    ws.row_dimensions[r].height = 18

# ---------- 채점 ----------
def grade_of(score, cuts):
    for i, c in enumerate(cuts):
        if score >= c: return i + 1
    return 0

def score_one(ans, key, point):
    got = [point[i] if i < len(ans) and ans[i] == key[i] else 0 for i in range(len(key))]
    return sum(got), got

# ---------- 성적표 한 장 ----------
BLOCK_H = 41
TC_, UC_ = 20, 21      # 그래프가 쓰는 칸 (인쇄 범위 밖)
LOGO = "로고-성적표.png"

def build_report(ws, top, st, avg, rnd):
    ws.row_dimensions[top].height = 40
    img = XLImage(LOGO); img.anchor = f"A{top}"; ws.add_image(img)
    put(ws, top+1, 1, f"  {rnd.title} 성적표", F(13, True, TITLE_INK), fill=BAND, align=L)
    merge(ws, top+1, 1, top+1, 16)
    ws.row_dimensions[top+1].height = 28
    ws.row_dimensions[top+2].height = 17

    # 1. 학생정보 — 선택과목이 있는 회차는 칸을 하나 더 둡니다
    tab(ws, top+3, 1, 3, "1. 학생정보")
    if rnd.choices:
        info = [("성명", 1, 4), ("선택과목", 5, 8), ("학년", 9, 11), ("시행일", 12, 16)]
        vals = [st["name"], st.get("choice", ""), rnd.who, rnd.date]
    else:
        info = [("성명", 1, 5), ("학년", 6, 10), ("시행일", 11, 16)]
        vals = [st["name"], rnd.who, rnd.date]
    head(ws, top+4, info)
    for (t, c1, c2), v in zip(info, vals):
        cell = put(ws, top+5, c1, v, F(11, True, VAL), fill=TD, align=C, border=SOFT)
        if t == "시행일": cell.number_format = "yyyy-mm-dd"
        if c2 > c1: merge(ws, top+5, c1, top+5, c2)
        for c in range(c1, c2+1): ws.cell(top+5, c).border = SOFT
    ws.row_dimensions[top+5].height = 24
    ws.row_dimensions[top+6].height = 17

    # 2. 성적
    tab(ws, top+7, 1, 3, "2. 성적")
    sc = [("점수 (100점 만점)", 1, 4), ("등급", 5, 6), ("응시생 평균", 7, 9), ("등급컷", 10, 16)]
    head(ws, top+8, sc)
    g = f'{st["grade"]}등급' if st["grade"] else "등급 외"
    svals = [(st["score"], "0", F(24, True, PURPLE)),
             (g, "General", F(11, True, VAL)),
             (avg, "0.0", F(11, True, VAL)),
             ("-".join(str(c) for c in st["cuts"]), "General", F(11, True, VAL))]
    for (t, c1, c2), (v, fmt, fnt) in zip(sc, svals):
        put(ws, top+9, c1, v, fnt, fill=TD, align=C, border=SOFT, fmt=fmt)
        if c2 > c1: merge(ws, top+9, c1, top+9, c2)
        for c in range(c1, c2+1): ws.cell(top+9, c).border = SOFT
    ws.row_dimensions[top+9].height = 32
    ws.row_dimensions[top+10].height = 17

    # 3. 영역분류별 성취도
    tab(ws, top+11, 1, 5, "3. 영역분류별 성취도 분석")
    head(ws, top+12, [("분류",1,1), ("영역",2,3), ("배점",4,4), ("득점",5,5), ("성취도%",6,7)],
         redcol=5)
    last = None
    for k in range(AREA_ROWS):
        r = top + 13 + k
        ws.row_dimensions[r].height = 19
        if k >= len(st["areas"]):          # 남는 줄은 테두리만 두고 비웁니다
            for c in range(1, 8): ws.cell(r, c).border = SOFT
            merge(ws, r, 2, r, 3); merge(ws, r, 6, r, 7)
            continue
        grp, name, a, b = st["areas"][k]
        full = sum(st["point"][a-1:b]); mine = sum(st["got"][a-1:b])
        put(ws, r, 1, "" if grp == last else grp, F(11, True, VAL), align=C, border=SOFT)
        last = grp
        put(ws, r, 2, name, F(11, True, VAL), align=C, border=SOFT)
        merge(ws, r, 2, r, 3); ws.cell(r, 3).border = SOFT
        put(ws, r, 4, full, F(11, True, VAL), align=C, border=SOFT)
        put(ws, r, 5, mine, F(11, True, VAL), align=C, border=SOFT)
        put(ws, r, 6, mine/full*100, F(11, True, VAL), align=C, border=SOFT, fmt="0.0")
        merge(ws, r, 6, r, 7); ws.cell(r, 7).border = SOFT
        put(ws, r, TC_, name, F(9, False, GRAY))
        put(ws, r, UC_, mine/full*100, F(9, False, GRAY), fmt="0.0")
    rs = top + 24
    topline = Border(left=_edge, right=_edge, bottom=_edge,
                     top=Side(style="medium", color="9CA3AF"))
    put(ws, rs, 1, "합계", F(11, True, VAL), fill=SUM_, align=C, border=topline)
    merge(ws, rs, 1, rs, 3)
    for c in (2,3): ws.cell(rs, c).border = topline
    put(ws, rs, 4, 100, F(11, True, VAL), fill=SUM_, align=C, border=topline)
    put(ws, rs, 5, st["score"], F(11, True, VAL), fill=SUM_, align=C, border=topline)
    put(ws, rs, 6, st["score"], F(11, True, VAL), fill=SUM_, align=C, border=topline, fmt="0.0")
    merge(ws, rs, 6, rs, 7); ws.cell(rs, 7).border = topline
    ws.row_dimensions[rs].height = 20

    # 막대그래프
    ch = BarChart(); ch.type = "col"; ch.title = "영역별 성취도 (%)"
    try:
        ch.title.tx.rich.p[0].pPr = ParagraphProperties(
            defRPr=CharacterProperties(sz=1100, b=True, solidFill=TITLE_INK, latin=None))
    except Exception: pass
    ch.legend = None; ch.gapWidth = 55; ch.height, ch.width = 8.2, 9.3
    ch.add_data(Reference(ws, min_col=UC_, min_row=top+13, max_row=top+23), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=TC_, min_row=top+13, max_row=top+23))
    ser = ch.series[0]
    ser.graphicalProperties.solidFill = PURPLE
    ser.graphicalProperties.line.noFill = True
    ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100; ch.y_axis.majorUnit = 20
    ch.y_axis.delete = False; ch.x_axis.delete = False; ch.dispBlanksAs = "gap"
    def _axfont(ax, pt):
        ax.txPr = RichText(bodyPr=RichTextProperties(),
                           p=[Paragraph(pPr=ParagraphProperties(
                               defRPr=CharacterProperties(sz=pt, latin=None)))])
    _axfont(ch.x_axis, 800); _axfont(ch.y_axis, 850)
    ws.add_chart(ch, f"H{top+12}")
    ws.row_dimensions[top+25].height = 17

    # 4. 문항 채점표
    tab(ws, top+26, 1, 3, "4. 문항 채점표")
    for b in range(3):
        r0 = top + 27 + 4*b
        for j, t in enumerate(["문항 번호", "정답", "학생답안", "정오"]):
            put(ws, r0+j, 1, t, F(11, True, THINK), fill=TH, align=C, border=SOFT)
            ws.row_dimensions[r0+j].height = 16
        for j in range(15):
            q = b*15 + j + 1; c = 2 + j
            put(ws, r0,   c, q, F(11, True, VAL), fill=TH, align=C, border=SOFT)
            put(ws, r0+1, c, st["key"][q-1], F(11), align=C, border=SOFT)
            mine = st["ans"][q-1] if q-1 < len(st["ans"]) else 0
            put(ws, r0+2, c, mine if mine else "—", F(11), align=C, border=SOFT)
            put(ws, r0+3, c, st["got"][q-1], F(11, True, VAL), align=C, border=SOFT)
        ws.conditional_formatting.add(
            f"B{r0+3}:P{r0+3}",
            CellIsRule(operator="equal", formula=["0"], fill=FILL(MINT),
                       font=Font(name=FAM, size=11, bold=True, color=MINT_INK)))
    ws.row_dimensions[top+39].height = 1
    ws.row_dimensions[top+40].height = 1

def style_report_sheet(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 11.5
    for c in range(2, 17): ws.column_dimensions[CL(c)].width = 5.6
    for c in (TC_, UC_): ws.column_dimensions[CL(c)].width = 11
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.sheet_format.customHeight = True
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = ws.page_margins.right = 0.55
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.page_margins.header = ws.page_margins.footer = 0.2
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True

# ---------- 만들기 ----------
def main(round_title, students, out, exams_path):
    book = read_exams(exams_path)
    hit = [e for e in (book['exams'] or []) if e.get('title') == round_title]
    if not hit:
        names = ' / '.join(e.get('title','') for e in (book['exams'] or []))
        raise SystemExit('exams.js 에 "%s" 회차가 없습니다.\n있는 회차: %s' % (round_title, names))
    rnd = Round(hit[0], book)

    for st in students:
        ch = st.get('choice') or (rnd.choices[0] if rnd.choices else '')
        if rnd.choices and ch not in rnd.choices:
            raise SystemExit('%s 학생의 선택과목 "%s" 을 모르겠습니다.' % (st.get('name'), ch))
        st['choice'] = ch if rnd.choices else ''
        st['key']   = rnd.key(ch)
        st['point'] = rnd.point(ch)
        st['cuts']  = rnd.cuts(ch)
        st['areas'] = rnd.areas(ch)
        st['score'], st['got'] = score_one(st['ans'], st['key'], st['point'])
        st['grade'] = grade_of(st['score'], st['cuts'])
        if sum(st['point']) != 100:
            raise SystemExit('%s 의 배점 합이 %d점입니다.' % (ch or '공통', sum(st['point'])))
    avg = sum(s['score'] for s in students) / len(students)
    targets = rnd.choices if rnd.choices else ['']

    wb = Workbook(); wb.remove(wb.active)

    al = wb.create_sheet("성적표(전원)")
    for i, st in enumerate(students):
        top = 1 + BLOCK_H * i
        build_report(al, top, st, avg, rnd)
        if i: al.row_breaks.append(Break(id=top-1))
    style_report_sheet(al)
    al.page_setup.fitToHeight = 0
    al.print_area = f"A1:P{BLOCK_H*len(students)}"

    # ---------- 분석 ----------
    an = wb.create_sheet("분석")
    an.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [12, 14, 12, 12, 14, 14, 12, 6]):
        an.column_dimensions[col].width = w
    title = (rnd.who + ' · ' if rnd.who else '') + rnd.title + ' 분석'
    put(an, 1, 1, title, F(15, True, TITLE_INK), align=L); merge(an, 1, 1, 1, 6)
    put(an, 2, 1, str(rnd.date), F(10, False, GRAY), align=L)

    def heading(r, text):
        put(an, r, 1, text, F(12, True, PURPLE), align=L); merge(an, r, 1, r, 6)
        return r + 1
    def header(r, cols):
        for c, t in enumerate(cols):
            put(an, r, 1+c, t, F(11, True, THINK), fill=TH, align=C, border=SOFT)
        return r + 1
    def row(r, vals, bold_at=None):
        for c, v in enumerate(vals):
            f = F(11, True, VAL) if c == bold_at else F(11)
            cell = put(an, r, 1+c, v, f, align=C, border=SOFT)
            if isinstance(v, float): cell.number_format = "0.0"
        return r + 1

    r = heading(4, "■ 한눈에 보기")
    box = [("응시 인원", f"{len(students)}명"), ("평균 점수", f"{avg:.1f}점"),
           ("최고 점수", f"{max(s['score'] for s in students)}점"),
           ("최저 점수", f"{min(s['score'] for s in students)}점")]
    for c, (t, v) in enumerate(box):
        put(an, r, 1+c, t, F(11, True, THINK), fill=TH, align=C, border=SOFT)
        put(an, r+1, 1+c, v, F(13, True, VAL), fill=TD, align=C, border=SOFT)
    an.row_dimensions[r+1].height = 24
    r += 3

    r = heading(r, "■ 등급 분포")
    r = header(r, ["등급", "인원", "비율%"])
    top_grade = max((len(s['cuts']) for s in students), default=4)
    for g in list(range(1, top_grade + 1)) + [0]:
        cnt = sum(1 for s in students if s['grade'] == g)
        if cnt == 0 and g == 0: continue
        r = row(r, [f"{g}등급" if g else "등급 외", cnt, cnt/len(students)*100])
    r += 1

    r = heading(r, "■ 학생별 점수 (높은 순)")
    cols = ["이름"] + (["선택과목"] if rnd.choices else []) + ["점수", "등급", "틀린 문항 수", "가장 약한 영역"]
    r = header(r, cols)
    for s in sorted(students, key=lambda s: -s['score']):
        pct = [(sum(s['got'][a-1:b])/sum(s['point'][a-1:b]), nm) for _, nm, a, b in s['areas']]
        vals = [s['name']] + ([s['choice']] if rnd.choices else []) + [
            s['score'], f"{s['grade']}등급" if s['grade'] else "등급 외",
            sum(1 for q in range(len(s['key'])) if s['got'][q] == 0), min(pct)[1]]
        r = row(r, vals, bold_at=(2 if rnd.choices else 1))
    r += 1

    r = heading(r, "■ 영역별 평균 성취도")
    r = header(r, (["대상"] if rnd.choices else []) + ["분류", "영역", "배점", "평균 득점", "성취도%"])
    for ch in targets:
        mine = [s for s in students if s['choice'] == ch] if rnd.choices else students
        if not mine: continue
        for grp, nm, a, b in rnd.areas(ch):
            full = sum(mine[0]['point'][a-1:b])
            m = sum(sum(s['got'][a-1:b]) for s in mine) / len(mine)
            r = row(r, (([ch] if rnd.choices else []) + [grp, nm, full, m, m/full*100]))
    r += 1

    r = heading(r, "■ 문항별 정답률 (낮은 순)")
    r = header(r, (["대상"] if rnd.choices else []) + ["문항", "정답", "배점", "영역", "맞힌 인원", "정답률%"])
    common = rnd.exam.get('common')
    if common is None: common = book['common']
    if not rnd.choices: common = rnd.count
    rate = []
    for ch in targets:
        mine = [s for s in students if s['choice'] == ch] if rnd.choices else students
        if not mine: continue
        key, point = mine[0]['key'], mine[0]['point']
        areas = rnd.areas(ch)
        def area_of(q):
            for _, nm, a, b in areas:
                if a <= q <= b: return nm
            return ""
        first = 1 if ch == targets[0] else common + 1     # 공통은 한 번만 넣습니다
        for q in range(first, rnd.count + 1):
            ok = sum(1 for s in mine if q-1 < len(s['ans']) and s['ans'][q-1] == key[q-1])
            label = '공통' if q <= common else ch
            rate.append((ok/len(mine)*100, q, label, key[q-1], point[q-1], area_of(q), ok))
    for pctv, q, label, k, pt, ar, ok in sorted(rate, key=lambda t: (t[0], t[1])):
        vals = (([label] if rnd.choices else []) + [q, k, pt, ar, ok, pctv])
        rr = row(r, vals)
        if pctv < 40:
            cell = an.cell(r, len(vals))
            cell.fill = FILL(MINT); cell.font = Font(name=FAM, size=11, bold=True, color=MINT_INK)
        r = rr
    an.page_setup.orientation = "portrait"
    an.page_setup.paperSize = an.PAPERSIZE_A4

    # ---------- 답안 ----------
    aw = wb.create_sheet("답안")
    aw.sheet_view.showGridLines = False
    for col, w in zip("ABCD", [12, 12, 50, 8]):
        aw.column_dimensions[col].width = w
    put(aw, 1, 1, "채점에 쓴 답입니다 (1번부터 차례로 이어서 적었습니다)",
        F(12, True, TITLE_INK), align=L); merge(aw, 1, 1, 1, 4)
    r = header_row = 2
    cols = ["이름"] + (["선택과목"] if rnd.choices else []) + ["학생 답", "점수"]
    for c, t in enumerate(cols):
        put(aw, r, 1+c, t, F(11, True, THINK), fill=TH, align=C, border=SOFT)
    r += 1
    for s in students:
        vals = [s['name']] + ([s['choice']] if rnd.choices else []) + [
            "".join(str(a) if a else "-" for a in s['ans']), s['score']]
        for c, v in enumerate(vals):
            put(aw, r, 1+c, v, F(11, True, VAL) if c == len(vals)-1 else F(11),
                align=L if c == len(vals)-2 else C, border=SOFT)
        r += 1
    r += 1
    for ch in targets:
        put(aw, r, 1, ("정답 · " + ch) if ch else "정답", F(11, True, VAL), align=C, border=SOFT)
        merge(aw, r, 1, r, len(cols)-2)
        put(aw, r, len(cols)-1, "".join(str(k) for k in rnd.key(ch)), F(11, True, VAL), align=L, border=SOFT)
        r += 1
        put(aw, r, 1, ("배점 · " + ch) if ch else "배점", F(11, True, VAL), align=C, border=SOFT)
        merge(aw, r, 1, r, len(cols)-2)
        put(aw, r, len(cols)-1, "".join(str(p) for p in rnd.point(ch)), F(11, True, VAL), align=L, border=SOFT)
        r += 1

    wb.save(out)
    print("만들었습니다:", out)

if __name__ == "__main__":
    if len(sys.argv) < 4:
        raise SystemExit('쓰는 법: python3 성적표-만들기.py "회차 이름" 학생답안.json 성적표.xlsx [exams.js 경로]')
    title = sys.argv[1]
    data = json.load(open(sys.argv[2], encoding="utf-8"))
    out = sys.argv[3]
    exams = sys.argv[4] if len(sys.argv) > 4 else os.path.join("..", "exams.js")
    main(title, data, out, exams)
