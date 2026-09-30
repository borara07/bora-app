# 예비고(중3) 학업진단평가 성적표 엑셀을 만드는 스크립트입니다.
# (선생님은 이 파일을 쓰지 않습니다. 만들어진 엑셀·PDF만 쓰시면 됩니다)
# -*- coding: utf-8 -*-
"""예비고 진단평가 성적표 만들기

   모의고사(고3)와 달리 선택과목이 없고 등급이 8단계여서 따로 만들었습니다.
   학생 답안은 OMR 답안지를 읽어 넣습니다."""
import datetime, os, sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.pagebreak import Break
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (ParagraphProperties, CharacterProperties,
                                   RichTextProperties, Paragraph)
from openpyxl.worksheet.properties import PageSetupProperties

# ---------- 회차 내용 ----------
TITLE = "9월 학업진단평가"
WHO   = "예비고(중3)"
DATE  = datetime.date(2026, 9, 20)
KEY   = [int(c) for c in '445125412251123513342441212544143553223533531']
POINT = [int(c) for c in '223223222232222222322222322232223223222232232']
CUTS  = [76, 68, 60, 52, 44, 37, 30, 23]
AREAS = [('화법과 작문','화법',1,3), ('화법과 작문','화법+작문',4,7), ('화법과 작문','작문',8,10),
         ('문법','문법',11,15),
         ('독서','사회',16,20), ('독서','인문',21,25), ('독서','기술',26,30),
         ('문학','현대시',31,33), ('문학','갈래복합',34,37),
         ('문학','현대소설',38,41), ('문학','고전소설',42,45)]
assert sum(POINT) == 100

# ---------- 색·글꼴 (30차 성적표와 같게) ----------
DEEP, PRI, LIGHT, CARD, GRAY = "5B21B6", "7C3AED", "EDE9FE", "F5F3FF", "6B7280"
BAND, TAB, TH, TD, SUM_, EDGE, INK, BAR = ("EDE9FE","DDD6FE","EDE9FE","FBFAFF",
                                           "E9E4FF","D8D0F5","3B0764","8064A2")
MINT, MINT_INK = "B8E6E0", "0F4F49"
FAM = "나눔스퀘어라운드"

def F(sz=11, b=False, c="1F2937"): return Font(name=FAM, size=sz, bold=b, color=c)
def FILL(c): return PatternFill("solid", fgColor=c)
_edge = Side(style="thin", color=EDGE)
SOFT  = Border(left=_edge, right=_edge, top=_edge, bottom=_edge)
C = Alignment(horizontal="center", vertical="center")
L = Alignment(horizontal="left",   vertical="center")
LW= Alignment(horizontal="left",   vertical="center", wrap_text=True)

def put(ws, r, c, v, font=None, fill=None, align=None, border=None, fmt=None):
    cell = ws.cell(r, c, v)
    if font: cell.font = font
    if fill: cell.fill = FILL(fill)
    if align: cell.alignment = align
    if border: cell.border = border
    if fmt: cell.number_format = fmt
    return cell

def merge(ws, r1, c1, r2, c2): ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

def tab(ws, r, c1, c2, text):
    put(ws, r, c1, text, F(12, True, INK), fill=TAB, align=L)
    merge(ws, r, c1, r, c2); ws.row_dimensions[r].height = 21

def head(ws, r, groups, redcol=None):
    for t, c1, c2 in groups:
        cell = put(ws, r, c1, t, F(11, True, INK), fill=TH, align=C, border=SOFT)
        if redcol is not None and c1 == redcol: cell.fill = FILL(TAB)
        if c2 > c1: merge(ws, r, c1, r, c2)
        for c in range(c1, c2 + 1): ws.cell(r, c).border = SOFT
    ws.row_dimensions[r].height = 18

# ---------- 채점 ----------
def grade_of(s):
    for i, c in enumerate(CUTS):
        if s >= c: return i + 1
    return 0

def score_one(ans):
    got = [POINT[i] if ans[i] == KEY[i] else 0 for i in range(45)]
    return sum(got), got

# ---------- 성적표 한 장 ----------
BLOCK_H = 41
TC_, UC_ = 20, 21      # 그래프가 쓰는 칸 (인쇄 범위 밖)
LOGO = "로고-성적표.png"

def build_report(ws, top, st, avg):
    ws.row_dimensions[top].height = 40
    img = XLImage(LOGO); img.anchor = f"A{top}"; ws.add_image(img)
    put(ws, top+1, 1, f"{TITLE} 성적표", F(16, True, INK), fill=BAND, align=L)
    merge(ws, top+1, 1, top+1, 16)
    ws.row_dimensions[top+1].height = 26
    ws.row_dimensions[top+2].height = 17

    # 1. 학생정보
    tab(ws, top+3, 1, 3, "1. 학생정보")
    info = [("성명", 1, 5), ("학년", 6, 10), ("시행일", 11, 16)]
    head(ws, top+4, info)
    vals = [st["name"], WHO, DATE]
    for (t, c1, c2), v in zip(info, vals):
        cell = put(ws, top+5, c1, v, F(11, True, INK), fill=TD, align=C, border=SOFT)
        if t == "시행일": cell.number_format = "yyyy-mm-dd"
        if c2 > c1: merge(ws, top+5, c1, top+5, c2)
        for c in range(c1, c2+1): ws.cell(top+5, c).border = SOFT
    ws.row_dimensions[top+5].height = 24
    ws.row_dimensions[top+6].height = 17

    # 2. 성적
    tab(ws, top+7, 1, 3, "2. 성적")
    sc = [("점수 (100점 만점)", 1, 4), ("등급", 5, 6), ("전체 평균", 7, 9), ("등급컷", 10, 16)]
    head(ws, top+8, sc)
    g = f'{st["grade"]}등급' if st["grade"] else "등급 외"
    svals = [(st["score"], "0", F(24, True, PRI)),
             (g, "General", F(11, True, INK)),
             (avg, "0.0", F(11, True, INK)),
             ("-".join(str(c) for c in CUTS), "General", F(11, True, INK))]
    for (t, c1, c2), (v, fmt, fnt) in zip(sc, svals):
        put(ws, top+9, c1, v, fnt, fill=TD, align=C, border=SOFT, fmt=fmt)
        if c2 > c1: merge(ws, top+9, c1, top+9, c2)
        for c in range(c1, c2+1): ws.cell(top+9, c).border = SOFT
    ws.row_dimensions[top+9].height = 32
    ws.row_dimensions[top+10].height = 17

    # 3. 영역분류별 성취도
    tab(ws, top+11, 1, 5, "3. 영역분류별 성취도 분석")
    head(ws, top+12, [("분류",1,1), ("영역",2,3), ("배점",4,4), ("득점",5,5), ("성취도%",6,7)],
         redcol=5)
    last = None
    for k, (grp, name, a, b) in enumerate(AREAS):
        r = top + 13 + k
        full = sum(POINT[a-1:b]); mine = sum(st["got"][a-1:b])
        put(ws, r, 1, "" if grp == last else grp, F(11, True, INK), fill=TD, align=C, border=SOFT)
        last = grp
        put(ws, r, 2, name, F(11, True, INK), align=C, border=SOFT)
        merge(ws, r, 2, r, 3); ws.cell(r, 3).border = SOFT
        put(ws, r, 4, full, F(11, True, INK), align=C, border=SOFT)
        put(ws, r, 5, mine, F(11, True, INK), fill=TD, align=C, border=SOFT)
        put(ws, r, 6, mine/full*100, F(11, True, PRI), align=C, border=SOFT, fmt="0.0")
        merge(ws, r, 6, r, 7); ws.cell(r, 7).border = SOFT
        put(ws, r, TC_, name, F(9, False, GRAY))
        put(ws, r, UC_, mine/full*100, F(9, False, GRAY), fmt="0.0")
        ws.row_dimensions[r].height = 19
    rs = top + 24
    topline = Border(left=_edge, right=_edge, bottom=_edge,
                     top=Side(style="medium", color="A78BFA"))
    put(ws, rs, 1, "합계", F(11, True, INK), fill=SUM_, align=C, border=topline)
    merge(ws, rs, 1, rs, 3)
    for c in (2,3): ws.cell(rs, c).border = topline
    put(ws, rs, 4, 100, F(11, True, INK), fill=SUM_, align=C, border=topline)
    put(ws, rs, 5, st["score"], F(11, True, INK), fill=SUM_, align=C, border=topline)
    put(ws, rs, 6, st["score"], F(11, True, PRI), fill=SUM_, align=C, border=topline, fmt="0.0")
    merge(ws, rs, 6, rs, 7); ws.cell(rs, 7).border = topline
    ws.row_dimensions[rs].height = 20

    # 막대그래프
    ch = BarChart(); ch.type = "col"; ch.title = "영역별 성취도 (%)"
    try:
        ch.title.tx.rich.p[0].pPr = ParagraphProperties(
            defRPr=CharacterProperties(sz=1100, b=True, solidFill=INK, latin=None))
    except Exception: pass
    ch.legend = None; ch.gapWidth = 55; ch.height, ch.width = 8.2, 9.3
    ch.add_data(Reference(ws, min_col=UC_, min_row=top+13, max_row=top+23), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=TC_, min_row=top+13, max_row=top+23))
    ser = ch.series[0]
    ser.graphicalProperties.solidFill = BAR
    ser.graphicalProperties.line.noFill = True
    ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100; ch.y_axis.majorUnit = 20
    ch.y_axis.delete = False; ch.x_axis.delete = False; ch.dispBlanksAs = "gap"
    def _axfont(ax, pt):
        ax.txPr = RichText(bodyPr=RichTextProperties(),
                           p=[Paragraph(pPr=ParagraphProperties(
                               defRPr=CharacterProperties(sz=pt, latin=None)))])
    _axfont(ch.x_axis, 800); _axfont(ch.y_axis, 850)
    ws.add_chart(ch, f"H{top+12}")
    ws.row_dimensions[top+25].height = 17

    # 4. 문항 채점표
    tab(ws, top+26, 1, 3, "4. 문항 채점표")
    for b in range(3):
        r0 = top + 27 + 4*b
        for j, t in enumerate(["문항 번호", "정답", "학생답안", "정오"]):
            put(ws, r0+j, 1, t, F(11, True, INK), fill=TH if j == 0 else TD, align=C, border=SOFT)
            ws.row_dimensions[r0+j].height = 16
        for j in range(15):
            q = b*15 + j + 1; c = 2 + j
            put(ws, r0,   c, q, F(11, True, INK), fill=TH, align=C, border=SOFT)
            put(ws, r0+1, c, KEY[q-1], F(11), align=C, border=SOFT)
            put(ws, r0+2, c, st["ans"][q-1] if st["ans"][q-1] else "—", F(11), align=C, border=SOFT)
            put(ws, r0+3, c, st["got"][q-1], F(11, True, INK), align=C, border=SOFT)
        ws.conditional_formatting.add(
            f"B{r0+3}:P{r0+3}",
            CellIsRule(operator="equal", formula=["0"], fill=FILL(MINT),
                       font=Font(name=FAM, size=11, bold=True, color=MINT_INK)))
    ws.row_dimensions[top+39].height = 1
    ws.row_dimensions[top+40].height = 1

def style_report_sheet(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 11.5
    for c in range(2, 17): ws.column_dimensions[CL(c)].width = 5.6
    for c in (TC_, UC_): ws.column_dimensions[CL(c)].width = 11
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.sheet_format.customHeight = True
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = ws.page_margins.right = 0.55
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.page_margins.header = ws.page_margins.footer = 0.2
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True

# ---------- 만들기 ----------
def main(students, out):
    for st in students:
        st["score"], st["got"] = score_one(st["ans"])
        st["grade"] = grade_of(st["score"])
    avg = sum(s["score"] for s in students) / len(students)

    wb = Workbook(); wb.remove(wb.active)

    al = wb.create_sheet("성적표(전원)")
    for i, st in enumerate(students):
        top = 1 + BLOCK_H * i
        build_report(al, top, st, avg)
        if i: al.row_breaks.append(Break(id=top-1))
    style_report_sheet(al)
    al.page_setup.fitToHeight = 0
    al.print_area = f"A1:P{BLOCK_H*len(students)}"

    # 분석
    an = wb.create_sheet("분석")
    an.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [14, 12, 12, 12, 14, 14, 26, 6]):
        an.column_dimensions[col].width = w
    put(an, 1, 1, f"{WHO} · {TITLE} 분석", F(16, True, DEEP), align=L); merge(an, 1, 1, 1, 6)
    put(an, 2, 1, str(DATE), F(10, False, GRAY), align=L)
    r = 4
    put(an, r, 1, "■ 한눈에 보기", F(12, True, PRI), align=L); merge(an, r, 1, r, 6); r += 1
    box = [("응시 인원", f"{len(students)}명"), ("평균 점수", f"{avg:.1f}점"),
           ("최고 점수", f"{max(s['score'] for s in students)}점"),
           ("최저 점수", f"{min(s['score'] for s in students)}점")]
    for c, (t, v) in enumerate(box):
        put(an, r, 1+c, t, F(11, True, INK), fill=TH, align=C, border=SOFT)
        put(an, r+1, 1+c, v, F(13, True, PRI), fill=TD, align=C, border=SOFT)
    an.row_dimensions[r+1].height = 24
    r += 3
    put(an, r, 1, "■ 등급 분포", F(12, True, PRI), align=L); merge(an, r, 1, r, 6); r += 1
    for c, t in enumerate(["등급", "인원", "비율"]):
        put(an, r, 1+c, t, F(11, True, "FFFFFF"), fill=PRI, align=C, border=SOFT)
    r += 1
    for gnum in list(range(1, 9)) + [0]:
        cnt = sum(1 for s in students if s["grade"] == gnum)
        if cnt == 0 and gnum == 0: continue
        put(an, r, 1, f"{gnum}등급" if gnum else "등급 외", F(11), align=C, border=SOFT)
        put(an, r, 2, cnt, F(11), align=C, border=SOFT)
        put(an, r, 3, cnt/len(students)*100, F(11), align=C, border=SOFT, fmt="0.0")
        r += 1
    r += 1
    put(an, r, 1, "■ 학생별 점수 (높은 순)", F(12, True, PRI), align=L); merge(an, r, 1, r, 6); r += 1
    for c, t in enumerate(["이름", "점수", "등급", "틀린 문항 수", "가장 약한 영역"]):
        put(an, r, 1+c, t, F(11, True, "FFFFFF"), fill=PRI, align=C, border=SOFT)
    r += 1
    for s in sorted(students, key=lambda s: -s["score"]):
        pct = [(sum(s["got"][a-1:b])/sum(POINT[a-1:b]), nm) for _, nm, a, b in AREAS]
        weak = min(pct)[1]
        put(an, r, 1, s["name"], F(11), align=C, border=SOFT)
        put(an, r, 2, s["score"], F(11, True, INK), align=C, border=SOFT)
        put(an, r, 3, f'{s["grade"]}등급' if s["grade"] else "등급 외", F(11), align=C, border=SOFT)
        put(an, r, 4, sum(1 for q in range(45) if s["got"][q] == 0), F(11), align=C, border=SOFT)
        put(an, r, 5, weak, F(11), align=C, border=SOFT)
        r += 1
    r += 1
    put(an, r, 1, "■ 영역별 반 평균 성취도", F(12, True, PRI), align=L); merge(an, r, 1, r, 6); r += 1
    for c, t in enumerate(["분류", "영역", "배점", "반 평균 득점", "성취도%"]):
        put(an, r, 1+c, t, F(11, True, "FFFFFF"), fill=PRI, align=C, border=SOFT)
    r += 1
    for grp, nm, a, b in AREAS:
        full = sum(POINT[a-1:b])
        m = sum(sum(s["got"][a-1:b]) for s in students)/len(students)
        put(an, r, 1, grp, F(11), align=C, border=SOFT)
        put(an, r, 2, nm, F(11), align=C, border=SOFT)
        put(an, r, 3, full, F(11), align=C, border=SOFT)
        put(an, r, 4, m, F(11), align=C, border=SOFT, fmt="0.0")
        put(an, r, 5, m/full*100, F(11, True, PRI), align=C, border=SOFT, fmt="0.0")
        r += 1
    r += 1
    put(an, r, 1, "■ 문항별 정답률 (낮은 순)", F(12, True, PRI), align=L); merge(an, r, 1, r, 6); r += 1
    for c, t in enumerate(["문항", "정답", "배점", "영역", "맞힌 인원", "정답률%"]):
        put(an, r, 1+c, t, F(11, True, "FFFFFF"), fill=PRI, align=C, border=SOFT)
    r += 1
    def area_of(q):
        for _, nm, a, b in AREAS:
            if a <= q <= b: return nm
        return ""
    rate = []
    for q in range(1, 46):
        ok = sum(1 for s in students if s["ans"][q-1] == KEY[q-1])
        rate.append((ok/len(students)*100, q, ok))
    for pctv, q, ok in sorted(rate):
        put(an, r, 1, q, F(11), align=C, border=SOFT)
        put(an, r, 2, KEY[q-1], F(11), align=C, border=SOFT)
        put(an, r, 3, POINT[q-1], F(11), align=C, border=SOFT)
        put(an, r, 4, area_of(q), F(11), align=C, border=SOFT)
        put(an, r, 5, ok, F(11), align=C, border=SOFT)
        cell = put(an, r, 6, pctv, F(11, True, INK), align=C, border=SOFT, fmt="0.0")
        if pctv < 40: cell.fill = FILL(MINT); cell.font = Font(name=FAM, size=11, bold=True, color=MINT_INK)
        r += 1
    an.page_setup.orientation = "portrait"
    an.page_setup.paperSize = an.PAPERSIZE_A4

    # 답안 (선생님이 확인·수정할 수 있게 원본 답을 남깁니다)
    aw = wb.create_sheet("답안")
    aw.sheet_view.showGridLines = False
    aw.column_dimensions["A"].width = 6; aw.column_dimensions["B"].width = 12
    aw.column_dimensions["C"].width = 50; aw.column_dimensions["D"].width = 8
    put(aw, 1, 1, "OMR 답안지에서 읽어 낸 답입니다 (1~45번을 이어서 적었습니다)",
        F(12, True, DEEP), align=L); merge(aw, 1, 1, 1, 4)
    for c, t in enumerate(["답안지", "이름", "1~45번 답", "점수"]):
        put(aw, 2, 1+c, t, F(11, True, "FFFFFF"), fill=PRI, align=C, border=SOFT)
    for i, s in enumerate(students):
        put(aw, 3+i, 1, s["page"], F(11), align=C, border=SOFT)
        put(aw, 3+i, 2, s["name"], F(11), align=C, border=SOFT)
        put(aw, 3+i, 3, "".join(str(a) if a else "-" for a in s["ans"]), F(11), align=L, border=SOFT)
        put(aw, 3+i, 4, s["score"], F(11, True, INK), align=C, border=SOFT)
    put(aw, 4+len(students), 1, "정답", F(11, True, INK), align=C, border=SOFT)
    merge(aw, 4+len(students), 1, 4+len(students), 2)
    put(aw, 4+len(students), 3, "".join(str(k) for k in KEY), F(11, True, INK), align=L, border=SOFT)
    put(aw, 5+len(students), 1, "배점", F(11, True, INK), align=C, border=SOFT)
    merge(aw, 5+len(students), 1, 5+len(students), 2)
    put(aw, 5+len(students), 3, "".join(str(p) for p in POINT), F(11, True, INK), align=L, border=SOFT)

    wb.save(out)
    print("만들었습니다:", out)

if __name__ == "__main__":
    data = json.load(open(sys.argv[1], encoding="utf-8"))
    main(data, sys.argv[2] if len(sys.argv) > 2 else "성적표-예비고-9월진단평가.xlsx")
